from __future__ import annotations
from pathlib import Path
import json
from zipfile import ZipFile, ZIP_DEFLATED
from html import escape

def write_report(path: Path, title: str, payload: object) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        from docx import Document
        doc=Document(); doc.add_heading(title, 0); doc.add_paragraph("Automatically generated factual results.")
        doc.add_paragraph(json.dumps(payload, indent=2, default=str)); doc.save(path)
    except ImportError:
        body = f'<w:p><w:r><w:t>{escape(title)}</w:t></w:r></w:p><w:p><w:r><w:t>{escape(json.dumps(payload, default=str))}</w:t></w:r></w:p>'
        content = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'+body+'<w:sectPr/></w:body></w:document>'
        rels = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
        types = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
        with ZipFile(path, "w", ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", types); z.writestr("_rels/.rels", rels); z.writestr("word/document.xml", content)
    return path

def build_reports(output: Path, payloads: dict[str, object]) -> list[Path]:
    reports=output/"reports"; reports.mkdir(parents=True, exist_ok=True)
    names=["01_single_model_results.docx","02_agreement_method_comparison.docx","03_final_single_vs_multi_comparison.docx","04_COMPLETE_THESIS_RESULTS.docx"]
    return [write_report(reports/name, name.removesuffix(".docx"), payloads) for name in names]

def build_figures(output: Path, final: dict, agreements: list[dict]) -> list[Path]:
    try:
        import matplotlib.pyplot as plt
    except ImportError: raise RuntimeError("Figure creation requires matplotlib")
    figdir=output/"figures"; figdir.mkdir(parents=True, exist_ok=True)
    f,ax=plt.subplots(figsize=(8,5)); names=[x["method"] for x in agreements]; pos=list(range(len(names))); width=.38; ax.bar([i-width/2 for i in pos],[x.get("accuracy",0) for x in agreements],width,label="Accuracy"); ax.bar([i+width/2 for i in pos],[x.get("critical_safety_error_rate",0) for x in agreements],width,label="Critical Safety Error Rate"); ax.set_xticks(pos,names,rotation=30,ha="right"); ax.set_ylabel("Rate / score"); ax.set_ylim(0,1); ax.legend(); f.tight_layout(); p1=figdir/"figure_1_agreement_methods.png"; f.savefig(p1,dpi=300); plt.close(f)
    f,ax=plt.subplots(figsize=(7,5));
    for name,v in final.get("systems",{}).items(): ax.scatter(v.get("accuracy",0),v.get("critical_safety_error_rate",0),s=80); ax.annotate(name,(v.get("accuracy",0),v.get("critical_safety_error_rate",0)),xytext=(5,5),textcoords="offset points")
    ax.set_xlabel("Accuracy (higher is better)"); ax.set_ylabel("Critical Safety Error Rate (lower is better)"); ax.grid(alpha=.2); f.tight_layout(); p2=figdir/"figure_2_accuracy_vs_safety.png"; f.savefig(p2,dpi=300); plt.close(f)
    return [p1,p2]
