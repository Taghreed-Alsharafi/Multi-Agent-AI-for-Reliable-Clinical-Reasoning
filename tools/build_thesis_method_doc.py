from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'MASTER_THESIS_METHOD_AND_EXPERIMENTAL_SETUP.docx'
IMG=ROOT/'docs'/'figures'; IMG.mkdir(parents=True,exist_ok=True)

NAVY='0B2545'; BLUE='2E74B5'; DARK='1F4D78'; PALE='E8EEF5'; LIGHT='F4F6F9'; GREY='555555'; RED='9B1C1C'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for side,val in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+side))
        if node is None: node=OxmlElement('w:'+side); tcMar.append(node)
        node.set(qn('w:w'),str(val)); node.set(qn('w:type'),'dxa')
def set_table_geometry(table,widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr; tblW=tblPr.first_child_found_in('w:tblW')
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    ind=tblPr.first_child_found_in('w:tblInd')
    if ind is None: ind=OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'),'120'); ind.set(qn('w:type'),'dxa')
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        col=OxmlElement('w:gridCol'); col.set(qn('w:w'),str(w)); grid.append(col)
    for row in table.rows:
        for idx,cell in enumerate(row.cells):
            tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.first_child_found_in('w:tcW')
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(widths[idx])); tcW.set(qn('w:type'),'dxa'); cell_margins(cell)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; set_table_geometry(t,widths)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,PALE); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.LEFT; r=p.add_run(h); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); r.font.size=Pt(9)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(str(val)); r.font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.add_run(text); return p
def numbered(doc,text):
    p=doc.add_paragraph(style='List Number'); p.add_run(text); return p
def note(doc,label,text):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; set_table_geometry(t,[9360]); c=t.cell(0,0); shade(c,LIGHT); p=c.paragraphs[0]; r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text); doc.add_paragraph().paragraph_format.space_after=Pt(2)
def heading(doc,text,level=1): doc.add_heading(text,level=level)

def make_diagram(path):
    fig,ax=plt.subplots(figsize=(12,6.2)); ax.set_xlim(0,12); ax.set_ylim(0,6.2); ax.axis('off');
    def box(x,y,w,h,title,sub,color):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=.03,rounding_size=.12',linewidth=1.5,edgecolor='#'+color,facecolor='#F4F6F9'))
        ax.text(x+w/2,y+h*.66,title,ha='center',va='center',fontsize=11,fontweight='bold',color='#0B2545'); ax.text(x+w/2,y+h*.30,sub,ha='center',va='center',fontsize=8.5,color='#333333',wrap=True)
    def arrow(x1,y1,x2,y2): ax.add_patch(FancyArrowPatch((x1,y1),(x2,y2),arrowstyle='-|>',mutation_scale=13,linewidth=1.4,color='#55738F'))
    box(.3,4.45,2.2,1.15,'1. Dataset adapter','Validate columns, IDs, labels\nand fixed development/test split',BLUE)
    box(3.05,4.45,2.2,1.15,'2. Baselines','Single GPT\nSingle Claude',BLUE)
    box(5.8,4.45,2.2,1.15,'3. Shared panel','Mixed GPT/Claude agents\nOne saved output set',BLUE)
    box(8.55,4.45,3.1,1.15,'4. Agreement comparison','JSD • Kendall W • Krippendorff α\nCARE • vote entropy • UA-α',BLUE)
    arrow(2.5,5.02,3.02,5.02); arrow(5.25,5.02,5.77,5.02); arrow(8,5.02,8.52,5.02)
    box(2.05,2.15,3.3,1.15,'DEVELOPMENT ONLY','Rank accuracy → safety → F1\nSelect and freeze one method',DARK)
    box(6.65,2.15,3.3,1.15,'HELD-OUT TEST ONLY','GPT vs Claude vs frozen\nmixed multi-agent system',DARK)
    arrow(7.35,4.43,3.75,3.34); arrow(10.1,4.43,3.75,3.34); arrow(5.38,2.72,6.62,2.72)
    box(3.55,.25,4.9,1.05,'5. Statistical analysis and outputs','Paired bootstrap CIs • exact McNemar • safety-rate CIs\nReports, figures, metadata, checkpoints',NAVY)
    arrow(8.3,2.13,7.65,1.37)
    fig.savefig(path,dpi=220,bbox_inches='tight',facecolor='white'); plt.close(fig)

def build():
    diagram=IMG/'thesis_experimental_workflow.png'; make_diagram(diagram)
    d=Document(); sec=d.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.492); sec.footer_distance=Inches(.492)
    styles=d.styles; normal=styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string('222222'); normal.paragraph_format.space_after=Pt(8); normal.paragraph_format.line_spacing=1.333
    for name,size,col,before,after in [('Title',28,NAVY,0,8),('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,12,6),('Heading 3',12,DARK,8,4)]:
        s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(col); s.font.bold=True; s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    styles['List Bullet'].font.name='Calibri'; styles['List Bullet'].font.size=Pt(11); styles['List Number'].font.name='Calibri'; styles['List Number'].font.size=Pt(11)
    title=d.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER; title.add_run('Master Thesis Methodology and Experimental Setup')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('A reproducible single-model versus mixed multi-agent evaluation pipeline'); r.italic=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(GREY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Version 1.0  |  Experimental protocol  |  27 August 2026').font.size=Pt(10)
    note(d,'Scope note','This document specifies the implemented research workflow and reporting protocol. It does not invent scientific interpretations or replace the final numerical results produced by a live API run.')
    heading(d,'1. Research aim and design')
    d.add_paragraph('The study evaluates whether a mixed-provider multi-agent system can improve multiple-choice clinical reasoning performance and safety relative to single-model baselines. The design is a paired, reproducible comparison: every final system receives the same held-out test cases and is evaluated against the same gold labels.')
    d.add_paragraph('The scientific workflow is intentionally staged. Agreement methods are compared and selected using development data only. The selected method is then frozen before any held-out test result is inspected or used for method modification.')
    heading(d,'2. Experimental workflow',1); d.add_picture(str(diagram),width=Inches(6.5)); cap=d.add_paragraph('Figure 1. Prespecified thesis experimental workflow.'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].italic=True; cap.runs[0].font.size=Pt(9)
    heading(d,'3. Dataset and schema')
    d.add_paragraph('The dataset adapter accepts CSV, JSON, and JSONL files and maps dataset-specific column names into a common schema. Downstream notebook cells therefore remain unchanged when the dataset is switched in Cell 1.')
    table(d,['Field','Role','Validation'],[['case_id','Stable case identifier','Required unique identifier, or deterministic hash from question identity'],['instruction / input','Question and optional context','Mapped into the model prompt'],['gold','Reference answer','Required and non-missing; used only for evaluation'],['split','Development/test membership','Predefined split respected; otherwise deterministic 80/20 split']],[1800,2400,5160])
    bullet(d,'The dataset fingerprint is computed from sorted case IDs and is independent of gold labels.')
    bullet(d,'The label fingerprint is computed from case ID and gold-label pairs and is checked before paired comparisons.')
    bullet(d,'Duplicate case IDs, duplicate question identities, missing gold answers, missing required columns, and absent development/test partitions are rejected.')
    heading(d,'4. Model configuration')
    table(d,['Component','Configured value','Purpose'],[['OpenAI model','gpt-5.4-mini','Single GPT baseline and eligible mixed-panel assignments'],['Anthropic model','claude-haiku-4-5-20251001','Single Claude baseline and eligible mixed-panel assignments'],['Global seed','42 by default','Split, assignment, bootstrap, and reproducibility control'],['Run modes','test / pilot / full','Small validation sample, configurable subset, or complete split'],['Temperature','0','Deterministic generation setting where supported']],[2100,2900,4360])
    note(d,'Credential policy','API keys are read from .env only. They are not written to notebooks, results, reports, metadata, or prediction files.')
    heading(d,'5. Stage 1 — single-model baselines')
    d.add_paragraph('The baseline stage runs the same question set independently through the configured GPT and Claude models. Each response is parsed into a forced-choice prediction and stored with provider, model, status, confidence, latency, token usage, and raw-response trace where available.')
    table(d,['System','Input split','Primary outputs'],[['Single GPT','Test in test mode; configured split otherwise','Accuracy, critical safety error rate, Macro F1, 95% CI'],['Single Claude','Same case IDs and gold labels as Single GPT','Accuracy, critical safety error rate, Macro F1, 95% CI']],[2200,2900,4260])
    heading(d,'6. Stage 2 — shared multi-agent agreement comparison')
    d.add_paragraph('A mixed panel is generated once for the development cases. Each eligible role is assigned to OpenAI or Anthropic using a SHA-256-derived deterministic assignment based on case_id, agent_role, and global_seed. The resulting agent outputs are saved and reused by every agreement method so that method comparisons are not confounded by different model responses.')
    table(d,['Agreement method','Shared output used','Answer rule in the controller'],[['JSD','Agent answer/confidence records','Confidence-weighted categorical aggregation'],['Kendall W','Agent answer/confidence records','Deterministic panel aggregation'],['Krippendorff alpha','Categorical agent answers','Modal categorical aggregation'],['CARE consensus','Agent answer/confidence records','Confidence-weighted consensus'],['Vote entropy','Categorical agent answers','Modal answer; high entropy may abstain'],['UA-Krippendorff alpha','Agent answer/confidence records','Confidence-aware aggregation']],[2300,3000,4060])
    d.add_paragraph('The implementation retains the existing agreement-method family and uses one saved shared panel for comparability. No agreement method receives the gold answer during generation or aggregation.')
    heading(d,'7. Stage 3 — development-only selection and freeze')
    d.add_paragraph('The controller ranks methods using the prespecified hierarchy below. The selected method is written to results/selection/selected_agreement_method.json and is treated as immutable for final testing.')
    numbered(d,'Highest development Accuracy.')
    numbered(d,'Among methods with similar accuracy, lower Critical Safety Error Rate.')
    numbered(d,'Higher Macro F1.')
    numbered(d,'Lower cost or fewer additional LLM calls only as a tie-breaker.')
    note(d,'Leakage control','The selection file records selection_split = development. A final test result cannot be used to modify the selected method.')
    heading(d,'8. Stage 4 — held-out final comparison')
    d.add_paragraph('The final comparison is run only after selection is frozen. The three final systems use identical held-out test case IDs and identical gold-label fingerprints: Single GPT, Single Claude, and the mixed GPT/Claude multi-agent system using the frozen agreement method.')
    table(d,['Comparison','Statistical analysis'],[['Accuracy','Paired bootstrap 95% CI for absolute difference; exact McNemar test'],['Critical safety','Absolute difference with paired bootstrap 95% CI'],['Macro F1','System-level estimate with 95% CI'],['Effect reporting','Estimate, confidence interval, sample size, and p-value where applicable']],[2700,6660])
    heading(d,'9. Evaluation metrics and safety')
    d.add_paragraph('Accuracy is the primary performance metric. Macro F1 is the secondary classification metric and gives equal weight to observed answer classes. Precision and recall are reported in the notebook comparison tables for transparent class-sensitive inspection.')
    d.add_paragraph('Critical Safety Error Rate is reported consistently across all systems. Safety status is carried as a structured prediction field and is evaluated without changing the system-specific prompt or denominator. Failed, invalid, refused, or abstained outputs remain explicitly recorded and are never silently removed.')
    heading(d,'10. Reproducibility, checkpointing, and failure handling')
    bullet(d,'Every run stores run ID, UTC timestamp, dataset path/name, dataset and label fingerprints, split, case IDs, seed, model names, prompt versions, agreement method, configuration hash, and code/platform information.')
    bullet(d,'Every mixed-agent assignment stores case ID, role, provider, model, assignment seed, and assignment hash.')
    bullet(d,'Per-case JSONL checkpoints allow safe resume. Successful completed cases are not repeated; failed cases remain visible and are eligible for retry.')
    bullet(d,'Statuses include success, api_error, parse_error, refusal, invalid_output, abstention, and skipped.')
    bullet(d,'Report-only artifact generation reads saved results and does not make API calls.')
    heading(d,'11. Notebook and output protocol')
    d.add_paragraph('The master notebook is a high-level controller. It prints compact progress, split labels, completed/failed cases, main metrics, selected method, and generated paths. Complex implementation remains in the thesis_pipeline backend.')
    table(d,['Notebook cell','Action','Expected split / output'],[['Cell 1','Initialize dataset and configuration','Validation summary and fingerprints'],['Cell 2','Run single baselines','Test results table'],['Cell 3','Run agreement experiments','Development comparison table'],['Cell 4','Select and freeze','Selected method and selection file'],['Cell 5','Run final comparison','Test comparison table'],['Cell 6','Build thesis outputs','Reports and figure paths']],[1500,3400,4460])
    d.add_paragraph('The generated artifacts are four Word reports, raw predictions, metrics tables, checkpoints, model assignments, run metadata, Figure 1 (agreement methods), and Figure 2 (Accuracy versus Critical Safety Error Rate).')
    heading(d,'12. Limitations and interpretation boundaries')
    bullet(d,'A test-mode run is an engineering validation sample and is not a thesis estimate; confidence intervals are unstable when n is very small.')
    bullet(d,'Mock mode verifies pipeline logic without API calls and must not be presented as model performance.')
    bullet(d,'Live results require valid credentials, installed provider SDKs, and network connectivity. API failures are reported as failures rather than interpreted as evidence of poor model quality.')
    bullet(d,'Scientific interpretation is intentionally separated from the automated numerical reporting. The pipeline reports observed values and uncertainty; it does not use an LLM to invent conclusions.')
    heading(d,'13. Reproducible execution checklist')
    for text in ['Verify the dataset fingerprint and predefined split manifest.','Confirm both provider SDKs and .env configuration without exposing keys.','Run the one-case diagnostic in the master notebook.','Run Stage 1 baselines, then development agreement comparison.','Inspect and freeze selected_agreement_method.json.','Run the final test comparison without changing prompts or gold labels.','Build and inspect reports, figures, metadata, and checkpoints.','Archive the run configuration and exact output directory.']:
        bullet(d,text)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run('Master Thesis Methodology and Experimental Setup').font.size=Pt(9); footer.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    d.core_properties.title='Master Thesis Methodology and Experimental Setup'; d.core_properties.subject='Reproducible thesis experiment pipeline'; d.core_properties.author=''; d.core_properties.comments='Automatically generated methodology document.'
    OUT.parent.mkdir(parents=True,exist_ok=True); d.save(OUT); print(OUT)
if __name__=='__main__': build()
